import os
import pytest
from docx import Document
from io import BytesIO

# 学生信息
STUDENT_INFO = {
    "name": "хао",
    "surname": "Цзя",
    "patronymic": "",
    "group": "4233k"
}

# 报告信息
REPORT_INFO = {
    "subject_name": "Основы программирования",
    "task_name": "Практическое задание №1",
    "task_type": "Практическое задание",
    "teacher": {
        "name": "Елена",
        "surname": "Шумова",
        "patronymic": "Олеговна",
        "status": "ПРЕПОДАВАТЕЛЬ"
    },
    "report_structure": ["Выход"],
    "uploaded_at": "2025-01-01T00:00:00Z"
}

@pytest.fixture
def specific_document_bytes():
    """返回特定文档的字节内容"""
    doc_path = os.path.join(os.path.dirname(__file__), "4233K_цзя хао_ЛР1.docx")
    if not os.path.exists(doc_path):
        pytest.skip(f"Specific document not found: {doc_path}")
    
    with open(doc_path, "rb") as f:
        return f.read()

def extract_full_text(docx_bytes):
    """提取文档中的所有文本（包括表格内容）"""
    doc = Document(BytesIO(docx_bytes))
    full_text = []
    
    # 添加段落文本
    for para in doc.paragraphs:
        full_text.append(para.text)
    
    # 添加表格文本
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text.append(cell.text)
    
    return "\n".join(full_text)

def test_document_has_required_sections(specific_document_bytes):
    """测试文档包含必要的章节"""
    full_text = extract_full_text(specific_document_bytes)
    
    # 检查必要的章节
    required_sections = REPORT_INFO["report_structure"]
    for section in required_sections:
        assert section in full_text, f"Required section '{section}' not found in document"

def test_document_has_title_page_info(specific_document_bytes):
    """测试文档标题页包含必要信息"""
    full_text = extract_full_text(specific_document_bytes)
    
    # 检查必要的信息 - 使用更灵活的方式
    required_info = [
        STUDENT_INFO["surname"],       # Цзя
        STUDENT_INFO["name"],           # хао
        STUDENT_INFO["group"],          # 4233k
        REPORT_INFO["subject_name"],    # Основы программирования
        REPORT_INFO["task_name"],       # Практическое задание №1
        REPORT_INFO["teacher"]["surname"],  # Шумова
        REPORT_INFO["teacher"]["name"],     # Елена
        "2025"  # 年份
    ]
    
    # 检查每个信息是否出现在文档的任何位置
    for info in required_info:
        assert info in full_text, f"Required info '{info}' not found in document"

def test_document_year(specific_document_bytes):
    """测试文档包含正确的年份"""
    full_text = extract_full_text(specific_document_bytes)
    assert "2025" in full_text, "Year 2025 not found in document"

def test_document_student_info(specific_document_bytes):
    """测试文档包含学生信息"""
    full_text = extract_full_text(specific_document_bytes)
    
    # 检查学生姓氏和名字是否出现在文档中
    assert STUDENT_INFO["surname"] in full_text, f"Student surname '{STUDENT_INFO['surname']}' not found"
    assert STUDENT_INFO["name"] in full_text, f"Student name '{STUDENT_INFO['name']}' not found"
    assert STUDENT_INFO["group"] in full_text, f"Student group '{STUDENT_INFO['group']}' not found"
