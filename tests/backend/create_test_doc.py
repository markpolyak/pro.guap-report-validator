from docx import Document

def create_valid_report_docx(file_path):
    """生成有效的报告文档"""
    doc = Document()
    
    # 文档内容
    doc.add_paragraph("МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ")
    doc.add_paragraph("Санкт-Петербургский государственный университет аэрокосмического приборостроения")
    doc.add_paragraph("Кафедра №43")
    doc.add_paragraph("Отчет по лабораторной работе №1")
    doc.add_paragraph("Студент: Иванов Иван Иванович Группа: 4931")
    doc.add_paragraph("Преподаватель: Антохина Юлия Анатольевна")
    doc.add_paragraph("Должность: Ректор, д.т.н., проф.")
    doc.add_paragraph("Цель работы: Изучение программирования")
    doc.add_paragraph("Задание: Реализовать алгоритм")
    doc.add_paragraph("Результат: Успешно")
    doc.add_paragraph("Выводы: Задание выполнено")
    doc.add_paragraph("Санкт-Петербург 2022")  # 确保这里的年份信息是正确的

    # 保存文档
    doc.save(file_path)

# 运行函数以创建文档
create_valid_report_docx('tests/backend/test_files/valid_report.docx')
