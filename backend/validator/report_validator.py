import re
from datetime import datetime
from dateutil import parser
from docx import Document
from docx.oxml.ns import qn
from io import BytesIO
import logging

logger = logging.getLogger(__name__)


class ReportValidator:
    def __init__(self, docx_bytes, student_info, report_info):
        self.docx_bytes = docx_bytes
        self.student_info = student_info
        self.report_info = report_info
        self.errors = []
        self.title_page_text = ""
        self.body_text = ""
        self.full_text = ""

    def _has_page_break(self, paragraph):
        """检查段落中是否包含分页符"""
        for run in paragraph.runs:
            for element in run._element:
                if element.tag.endswith('br') and element.get(qn('w:type')) == 'page':
                    return True
        return False

    def _extract_document(self):
        """解析文档并分离标题页和正文"""
        try:
            doc = Document(BytesIO(self.docx_bytes))

            if not doc.paragraphs:
                self.errors.append("Документ пуст")
                return

            # 提取整个文档文本
            self.full_text = "\n".join([para.text for para in doc.paragraphs])

            # 添加表格内容
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        self.full_text += "\n" + cell.text

            # 查找分页符位置
            page_break_index = None
            for i, para in enumerate(doc.paragraphs):
                if self._has_page_break(para):
                    page_break_index = i
                    break

            # 如果没有找到分页符，使用第一个一级标题作为分隔
            if page_break_index is None:
                for i, para in enumerate(doc.paragraphs):
                    if para.style and para.style.name.startswith('Heading 1') and i > 0:
                        page_break_index = i
                        break

            # 默认取前10段作为标题页
            if page_break_index is None:
                page_break_index = min(10, len(doc.paragraphs))

            # 分离标题页和正文
            self.title_page_text = "\n".join(
                para.text for para in doc.paragraphs[:page_break_index]
            )
            self.body_text = "\n".join(
                para.text for para in doc.paragraphs[page_break_index:]
            )

            # 添加表格内容到标题页
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        self.title_page_text += "\n" + cell.text.strip()

        except Exception as e:
            self.errors.append(f"Ошибка при разборе документа: {str(e)}")
            logger.exception("Ошибка при разборе документа")

    def _normalize_text(self, text):
        """规范化文本用于比较（小写+去多余空格）"""
        if not text:
            return ""
        # 保留标点符号，只去除多余空格
        return re.sub(r'\s+', ' ', text.lower().strip())

    def _check_title_page_field(self, field_label, field_value, required=True, is_year=False):
        """检查标题页特定字段"""
        if not field_value and required:
            self.errors.append(f"Отсутствует обязательное поле: {field_label}")
            return False

        if field_value:
            # 创建完整的字段文本（标签+值）
            if field_label:
                full_field_text = f"{field_label}: {field_value}"
            else:
                full_field_text = field_value

            normalized_field = self._normalize_text(full_field_text)
            normalized_title = self._normalize_text(self.title_page_text)

            # 对于年份，使用更宽松的匹配
            if is_year:
                # 查找所有4位数字序列
                years_in_text = re.findall(r'\d{4}', normalized_title)
                if field_value not in years_in_text:
                    self.errors.append(f"Не найден год выполнения отчета: {field_value}")
                    return False
                return True

            # 检查完整字段是否出现在标题页
            if normalized_field not in normalized_title:
                # 尝试部分匹配
                if any(part in normalized_title for part in normalized_field.split() if len(part) > 2):
                    return True

                self.errors.append(f"Не найдено на титульном листе: {full_field_text}")
                return False

        return True

    def _check_title_page(self):
        """验证标题页信息"""
        if not self.title_page_text:
            self.errors.append("Не удалось определить титульный лист")
            return

        # 组合学生全名
        student_fullname = " ".join(filter(None, [
            self.student_info.get("surname", ""),
            self.student_info.get("name", ""),
            self.student_info.get("patronymic", "")
        ]))

        # 组合教师全名
        teacher = self.report_info.get("teacher", {})
        teacher_fullname = " ".join(filter(None, [
            teacher.get("surname", ""),
            teacher.get("name", ""),
            teacher.get("patronymic", "")
        ]))

        # 提取年份
        uploaded_at = parser.parse(self.report_info["uploaded_at"])
        report_year = str(uploaded_at.year)

        # 检查字段 - 使用完整的标签+值组合
        required_fields = [
            ("тип задания", self.report_info.get("task_type", "")),
            ("студент", student_fullname),
            ("группа", self.student_info.get("group", "")),
            ("предмет", self.report_info.get("subject_name", "")),
            ("название задания", self.report_info.get("task_name", "")),
            ("преподаватель", teacher_fullname),
            ("должность", teacher.get("status", ""))
        ]

        # 特殊处理年份（通常没有标签前缀）
        self._check_title_page_field("", report_year, required=True, is_year=True)

        # 验证必填字段
        for label, value in required_fields:
            self._check_title_page_field(label, value, required=True)

    def _check_report_structure(self):
        """验证报告结构"""
        if not self.body_text:
            return

        sections = self.report_info.get("report_structure", [])
        if not sections:
            return

        normalized_body = self._normalize_text(self.body_text)

        for section in sections:
            if not section:
                continue

            # 检查章节标题是否在正文中
            normalized_section = self._normalize_text(section)
            if normalized_section not in normalized_body:
                # 尝试部分匹配
                keywords = normalized_section.split()
                if not any(keyword in normalized_body for keyword in keywords if len(keyword) > 2):
                    self.errors.append(f"Отсутствует раздел: {section}")

    def validate(self):
        """执行所有验证"""
        try:
            self._extract_document()
            if not self.errors:
                self._check_title_page()
                self._check_report_structure()
        except Exception as e:
            logger.exception("Ошибка при проверке отчета")
            self.errors.append(f"Системная ошибка: {str(e)}")

        return self.errors


if __name__ == '__main__':
    # 命令行使用示例
    import sys
    import json

    if len(sys.argv) != 4:
        print("Использование: python report_validator.py <docx_file> <student_info_json> <report_info_json>")
        sys.exit(1)

    docx_file = sys.argv[1]
    student_info = json.loads(sys.argv[2])
    report_info = json.loads(sys.argv[3])

    try:
        with open(docx_file, 'rb') as f:
            docx_bytes = f.read()

        validator = ReportValidator(docx_bytes, student_info, report_info)
        errors = validator.validate()

        if errors:
            print("Найдены ошибки:")
            for error in errors:
                print(f" - {error}")
        else:
            print("Отчет соответствует всем требованиям")
    except Exception as e:
        print(f"Ошибка: {str(e)}")
        sys.exit(1)
